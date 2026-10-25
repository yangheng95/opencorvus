"""Apply presentation-only revision against the first immutable delivery."""

from pathlib import Path
from io import BytesIO
import hashlib
import json
import subprocess
from zipfile import ZipFile
from lxml import etree
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.shared import Mm, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.table import Table
from docx.enum.table import WD_TABLE_ALIGNMENT

ROOT = Path(__file__).resolve().parents[2]
BASE = "5747fb6f85829b2c72edbba3101ee922ee508da2"


def original(path):
    return subprocess.check_output(["git", "show", f"{BASE}:{path}"], cwd=ROOT)


old_bytes = original("outputs/growth-plan.docx")
# This digest identifies independently published immutable attachment bytes,
# not an expectation recomputed from the revised document.
assert (
    hashlib.sha256(old_bytes).hexdigest()
    == "207561119225d920a84880620c7e05b60f443911195afa7fc4bc90610122060a"
)
md_reference = (
    ROOT
    / ".opencorvus/.r/tasks/tsk_g00VZ2GE9C00jkeX1GgT/artifacts/materializations/reads/2a03896811a36347f08c9eba2916c1e7e07b2695aa042eecf21bf8e8b2ddbc61-growth-plan.md"
)
assert (ROOT / "outputs/growth-plan.md").read_bytes() == md_reference.read_bytes()
old = Document(BytesIO(old_bytes))
doc = Document(BytesIO(old_bytes))
for sec in doc.sections:
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Mm(297), Mm(210)
    sec.top_margin = sec.bottom_margin = Mm(14)
for p in doc.paragraphs:
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.1
# Keep adjacent tables separate during import without adding any text.
separator = OxmlElement("w:p")
separator.set(qn("w:rsidR"), "00000001")
doc.tables[1]._tbl.addprevious(separator)
sp = Paragraph(separator, doc)
sp.paragraph_format.space_before = Pt(0)
sp.paragraph_format.space_after = Pt(0)
sp.paragraph_format.line_spacing = Pt(1)
widths = [[51, 42, 42, 42, 42, 42], [83, 83, 84], [80, 50, 131], [70, 191]]
for table, columns in zip(doc.tables, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table._tbl.tblPr.find(qn("w:tblW")).set(qn("w:type"), "dxa")
    table._tbl.tblPr.find(qn("w:tblW")).set(
        qn("w:w"), str(round(sum(columns) / 25.4 * 1440))
    )
    for col, width in zip(table.columns, columns):
        col.width = Mm(width)
    for index, row in enumerate(table.rows):
        prop = row._tr.get_or_add_trPr()
        if prop.find(qn("w:cantSplit")) is None:
            prop.append(OxmlElement("w:cantSplit"))
        if index == 0:
            prop.append(OxmlElement("w:tblHeader"))
        for cell, width in zip(row.cells, columns):
            cell.width = Mm(width)
            for p in cell.paragraphs:
                p.paragraph_format.keep_with_next = index == 0
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.space_after = Pt(3)
                for run in p.runs:
                    run.font.size = Pt(10.5)
out = ROOT / "outputs/growth-plan.docx"
doc.save(out)


def ordered(document):
    result = []
    for child in document.element.body:
        if child.tag == qn("w:p"):
            if child.get(qn("w:rsidR")) == "00000001":
                assert Paragraph(child, document).text == ""
                continue
            result.append({"paragraph": Paragraph(child, document).text})
        elif child.tag == qn("w:tbl"):
            result.append(
                {
                    "table": [
                        [[p.text for p in c.paragraphs] for c in r.cells]
                        for r in Table(child, document).rows
                    ]
                }
            )
    return result


new = Document(out)
assert ordered(old) == ordered(new), "Full ordered paragraphs/cells changed"
with ZipFile(out) as z:
    xml = etree.fromstring(z.read("word/document.xml"))
    styles = etree.fromstring(z.read("word/styles.xml"))
ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
for size in xml.findall(".//w:sectPr/w:pgSz", ns):
    assert size.get(qn("w:orient")) == "landscape"
    assert int(size.get(qn("w:w"))) > int(size.get(qn("w:h")))
for table in xml.findall(".//w:tbl", ns):
    rows = table.findall("w:tr", ns)
    assert rows[0].find("w:trPr/w:tblHeader", ns) is not None
    assert all(
        p.find("w:pPr/w:keepNext", ns) is not None
        for p in rows[0].findall("w:tc/w:p", ns)
    )
    assert all(r.find("w:trPr/w:cantSplit", ns) is not None for r in rows)
    grid = [int(c.get(qn("w:w"))) for c in table.findall("w:tblGrid/w:gridCol", ns)]
    assert sum(grid) <= round(261 / 25.4 * 1440) + 6
assert not xml.findall(".//w:pBdr", ns)
for name in ["Title", "Heading1", "Heading2"]:
    s = styles.xpath(f'//w:style[@w:styleId="{name}"]', namespaces=ns)[0]
    assert s.find("w:rPr/w:color", ns).get(qn("w:val")) == "000000"
    assert not s.findall(".//w:pBdr", ns)
report = {
    "status": "author_content_and_ooxml_checks_passed_pending_page_review",
    "baseline_artifact": "art_h28MoKiA1WMo4nQBkTLd revision 40",
    "baseline_snapshot": "19df3dd4-7067-562d-ac01-42acf9bfc56d",
    "baseline_git_blob_identity_verified_against_immutable_docx": True,
    "md_complete_bytes_equal_immutable_resource": True,
    "word_ordered_paragraphs_and_all_cell_paragraphs_equal": True,
    "added_empty_layout_separator_paragraphs": 1,
    "all_sections_landscape": True,
    "table_widths_mm": widths,
    "repeat_headers": 4,
    "header_keep_with_next": True,
    "all_22_rows_cantSplit": True,
    "heading_color": "000000",
    "decorative_paragraph_borders": 0,
    "content_note": "Original literal portrait/status wording is frozen; orientation change is presentation only.",
    "fact_check_items": [],
}
qa = ROOT / "qa/revised"
qa.mkdir(parents=True, exist_ok=True)
(qa / "content-comparison.json").write_text(
    json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
)
(qa / "complete-ordered-comparison.json").write_text(
    json.dumps(
        {"original": ordered(old), "revised": ordered(new), "equal": True},
        ensure_ascii=False,
        indent=2,
    ),
    encoding="utf-8",
)
print(json.dumps(report, ensure_ascii=False, indent=2))
