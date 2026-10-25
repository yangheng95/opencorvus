"""Generate matching local documents and compare their complete semantic sequences."""

from pathlib import Path
import json
import re
from html import escape
from html.parser import HTMLParser
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]


def md_blocks(text):
    blocks = []
    for line in text.splitlines():
        if not line.strip():
            continue
        if line.startswith("|"):
            cells = [c.strip() for c in line.strip("|").split("|")]
            if all(re.fullmatch(r":?-+:?", c) for c in cells):
                continue
            if blocks and blocks[-1][0] == "table":
                blocks[-1][1].append(cells)
            else:
                blocks.append(("table", [cells]))
        elif line.startswith("#"):
            level = len(line) - len(line.lstrip("#"))
            blocks.append((f"h{level}", line[level:].strip()))
        else:
            blocks.append(("p", line))
    return blocks


def set_font(run, size=11, bold=False):
    run.font.name = "Microsoft YaHei"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    run.bold = bold


def generate():
    text = (ROOT / "outputs/growth-plan.md").read_text(encoding="utf-8")
    blocks = md_blocks(text)
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.8)
    section.left_margin = section.right_margin = Cm(1.7)
    style = doc.styles["Normal"]
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.18
    parts = []
    for kind, body in blocks:
        if kind == "table":
            table = doc.add_table(rows=0, cols=len(body[0]))
            table.style = "Table Grid"
            table.autofit = False
            widths = (
                [2.35, 1.35, 1.2, 1.2, 1.7, 2.1, 1.8, 3.9]
                if len(body[0]) == 8
                else [3.8, 2.4, 11.4]
                if len(body[0]) == 3 and body[0][0] == "用途"
                else [3.4, 7.1, 7.1]
            )
            # Fit every table to the same portrait text width.
            widths = [w * 17.6 / sum(widths) for w in widths]
            parts.append("<table>")
            for i, row in enumerate(body):
                cells = table.add_row().cells
                trpr = table.rows[-1]._tr.get_or_add_trPr()
                trpr.append(OxmlElement("w:cantSplit"))
                if i == 0:
                    trpr.append(OxmlElement("w:tblHeader"))
                parts.append("<tr>")
                for j, value in enumerate(row):
                    cells[j].width = Cm(widths[j])
                    p = cells[j].paragraphs[0]
                    p.paragraph_format.space_after = Pt(3)
                    p.paragraph_format.line_spacing = 1.08
                    set_font(p.add_run(value), 9 if len(row) == 8 else 10, i == 0)
                    parts.append(f"<td>{escape(value)}</td>")
                parts.append("</tr>")
            parts.append("</table>")
        else:
            p = doc.add_paragraph(style="Normal")
            if kind.startswith("h"):
                p.style = doc.styles["Title" if kind == "h1" else "Heading 1"]
                p.paragraph_format.keep_with_next = True
                set_font(p.add_run(body), 19 if kind == "h1" else 14, True)
            else:
                set_font(p.add_run(body))
            parts.append(f"<{kind}>{escape(body)}</{kind}>")
    doc.save(ROOT / "outputs/growth-plan.docx")
    (ROOT / "artifacts/marketing-growth").mkdir(parents=True, exist_ok=True)
    (ROOT / "artifacts/marketing-growth/campaign-plan.md").write_text(
        text, encoding="utf-8"
    )
    html = (
        '<!doctype html><html lang="zh-CN"><meta charset="utf-8"><title>小型团队文档整理工具：本地营销学习计划</title><style>body{max-width:1000px;margin:40px auto;padding:0 24px;color:#000;font:17px/1.8 "Microsoft YaHei",sans-serif}table{border-collapse:collapse;width:100%;font-size:14px}td{border:1px solid #555;padding:8px;overflow-wrap:anywhere}tr:first-child{font-weight:bold}h1,h2{color:#000}@media(max-width:700px){body{padding:0 12px}table{font-size:12px}}</style><body>'
        + "".join(parts)
        + "</body></html>"
    )
    (ROOT / "outputs/growth-plan.html").write_text(html, encoding="utf-8")


class Reader(HTMLParser):
    def __init__(self):
        super().__init__()
        self.blocks, self.kind, self.value, self.rows, self.cells = [], None, "", [], []

    def handle_starttag(self, tag, attrs):
        if tag in ("h1", "h2", "p", "td"):
            self.kind, self.value = tag, ""
        elif tag == "table":
            self.rows = []
        elif tag == "tr":
            self.cells = []

    def handle_data(self, data):
        if self.kind:
            self.value += data

    def handle_endtag(self, tag):
        if tag in ("h1", "h2", "p"):
            self.blocks.append((tag, self.value))
            self.kind = None
        elif tag == "td":
            self.cells.append(self.value)
            self.kind = None
        elif tag == "tr":
            self.rows.append(self.cells)
        elif tag == "table":
            self.blocks.append(("table", self.rows))


def verify():
    text = (ROOT / "outputs/growth-plan.md").read_text(encoding="utf-8")
    expected = md_blocks(text)
    doc = Document(ROOT / "outputs/growth-plan.docx")
    actual = []
    for element in doc.element.body:
        if element.tag == qn("w:p"):
            from docx.text.paragraph import Paragraph

            p = Paragraph(element, doc)
            kind = (
                "h1"
                if p.style.name == "Title"
                else "h2"
                if p.style.name == "Heading 1"
                else "p"
            )
            actual.append((kind, p.text))
        elif element.tag == qn("w:tbl"):
            from docx.table import Table

            actual.append(
                ("table", [[c.text for c in r.cells] for r in Table(element, doc).rows])
            )
    reader = Reader()
    reader.feed((ROOT / "outputs/growth-plan.html").read_text(encoding="utf-8"))
    assert actual == expected, "Word semantic sequence differs"
    assert reader.blocks == expected, "HTML semantic sequence differs"
    assert (ROOT / "artifacts/marketing-growth/campaign-plan.md").read_text(
        encoding="utf-8"
    ) == text
    assert all(s.page_width < s.page_height for s in doc.sections)
    for p in doc.paragraphs:
        for r in p.runs:
            assert str(r.font.color.rgb) == "000000"
    result = {
        "method": "Parse MD headings/paragraphs/table rows; reopen DOCX body XML in document order and parse HTML independently. Exact ordered text and cell comparison; no digest shortcut.",
        "semantic_blocks": len(expected),
        "tables": sum(k == "table" for k, _ in expected),
        "table_cells": sum(
            sum(len(row) for row in b) for k, b in expected if k == "table"
        ),
        "chinese_characters": len(re.findall(r"[\u4e00-\u9fff]", text)),
        "md_docx_html_sequence_equal": True,
        "mirror_equal": True,
        "all_sections_portrait": True,
        "all_paragraph_runs_black": True,
    }
    (ROOT / "qa").mkdir(exist_ok=True)
    (ROOT / "qa/body-consistency.json").write_text(
        json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps(result, ensure_ascii=False))


if __name__ == "__main__":
    import sys

    if "--verify-only" not in sys.argv:
        generate()
    verify()
