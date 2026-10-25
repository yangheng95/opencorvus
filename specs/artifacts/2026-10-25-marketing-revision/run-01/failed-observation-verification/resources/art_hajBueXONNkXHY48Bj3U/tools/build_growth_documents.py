"""Generate matching local documents and compare their complete semantic sequences."""

from pathlib import Path
import json
import re
from html.parser import HTMLParser
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.section import WD_ORIENT
from io import BytesIO
import hashlib
import subprocess

ROOT = Path(__file__).resolve().parents[1]


def md_blocks(text):
    blocks = []
    for line in text.splitlines():
        if not line.strip():
            continue
        if line.startswith("|"):
            cells = [c.strip() for c in line.strip("|").split("|")]
            if all(re.fullmatch(r":?-+:?", c) for c in cells):
                continue
            if blocks and blocks[-1][0] == "table":
                blocks[-1][1].append(cells)
            else:
                blocks.append(("table", [cells]))
        elif line.startswith("#"):
            level = len(line) - len(line.lstrip("#"))
            blocks.append((f"h{level}", line[level:].strip()))
        else:
            blocks.append(("p", line))
    return blocks


def set_font(run, size=11, bold=False):
    run.font.name = "Microsoft YaHei"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    run.bold = bold


def generate():
    text = (ROOT / "outputs/growth-plan.md").read_text(encoding="utf-8")
    blocks = md_blocks(text)
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Cm(29.7), Cm(21)
    section.top_margin = section.bottom_margin = Cm(1.8)
    section.left_margin = section.right_margin = Cm(1.7)
    style = doc.styles["Normal"]
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.18
    # The template Title style supplies the blue bottom border, not its runs.
    for border in list(doc.styles["Title"].element.iter(qn("w:pBdr"))):
        border.getparent().remove(border)
    for kind, body in blocks:
        if kind == "table":
            table = doc.add_table(rows=0, cols=len(body[0]))
            table.style = "Table Grid"
            table.autofit = False
            widths = (
                [2.35, 1.35, 1.2, 1.2, 1.7, 2.1, 1.8, 3.9]
                if len(body[0]) == 8
                else [3.8, 2.4, 11.4]
                if len(body[0]) == 3 and body[0][0] == "用途"
                else [3.4, 7.1, 7.1]
            )
            widths = [w * 26.3 / sum(widths) for w in widths]
            for column, width in zip(table.columns, widths):
                column.width = Cm(width)
            for i, row in enumerate(body):
                cells = table.add_row().cells
                trpr = table.rows[-1]._tr.get_or_add_trPr()
                trpr.append(OxmlElement("w:cantSplit"))
                if i == 0:
                    trpr.append(OxmlElement("w:tblHeader"))
                for j, value in enumerate(row):
                    cells[j].width = Cm(widths[j])
                    p = cells[j].paragraphs[0]
                    p.paragraph_format.space_after = Pt(3)
                    p.paragraph_format.line_spacing = 1.08
                    p.paragraph_format.keep_with_next = i == 0
                    set_font(p.add_run(value), 10, i == 0)
        else:
            p = doc.add_paragraph(style="Normal")
            if kind.startswith("h"):
                p.style = doc.styles["Title" if kind == "h1" else "Heading 1"]
                p.paragraph_format.keep_with_next = True
                set_font(p.add_run(body), 19 if kind == "h1" else 14, True)
            else:
                set_font(p.add_run(body))
    doc.save(ROOT / "outputs/growth-plan.docx")


class Reader(HTMLParser):
    def __init__(self):
        super().__init__()
        self.blocks, self.kind, self.value, self.rows, self.cells = [], None, "", [], []

    def handle_starttag(self, tag, attrs):
        if tag in ("h1", "h2", "p", "td"):
            self.kind, self.value = tag, ""
        elif tag == "table":
            self.rows = []
        elif tag == "tr":
            self.cells = []

    def handle_data(self, data):
        if self.kind:
            self.value += data

    def handle_endtag(self, tag):
        if tag in ("h1", "h2", "p"):
            self.blocks.append((tag, self.value))
            self.kind = None
        elif tag == "td":
            self.cells.append(self.value)
            self.kind = None
        elif tag == "tr":
            self.rows.append(self.cells)
        elif tag == "table":
            self.blocks.append(("table", self.rows))


def verify():
    text = (ROOT / "outputs/growth-plan.md").read_text(encoding="utf-8")
    expected = md_blocks(text)
    doc = Document(ROOT / "outputs/growth-plan.docx")
    actual = []
    for element in doc.element.body:
        if element.tag == qn("w:p"):
            from docx.text.paragraph import Paragraph

            p = Paragraph(element, doc)
            kind = (
                "h1"
                if p.style.name == "Title"
                else "h2"
                if p.style.name == "Heading 1"
                else "p"
            )
            actual.append((kind, p.text))
        elif element.tag == qn("w:tbl"):
            from docx.table import Table

            actual.append(
                ("table", [[c.text for c in r.cells] for r in Table(element, doc).rows])
            )
    reader = Reader()
    reader.feed((ROOT / "outputs/growth-plan.html").read_text(encoding="utf-8"))
    assert actual == expected, "Word semantic sequence differs"
    assert reader.blocks == expected, "HTML semantic sequence differs"
    assert (ROOT / "artifacts/marketing-growth/campaign-plan.md").read_text(
        encoding="utf-8"
    ) == text
    assert all(s.page_width > s.page_height for s in doc.sections)
    baseline_md = (
        ROOT
        / ".opencorvus/.r/tasks/tsk_g00VZ1u3kk00D81zMYvb/artifacts/materializations/reads/fd9682e65a7fe3dd7369a536343d193d302138dc2e480604812025993e3ebb6b-growth-plan.md"
    )
    assert baseline_md.read_bytes() == (ROOT / "outputs/growth-plan.md").read_bytes()
    baseline_bytes = subprocess.check_output(
        ["git", "show", "1c388fe:outputs/growth-plan.docx"], cwd=ROOT
    )
    # Authenticate the Git object against the independently published first resource.
    assert (
        hashlib.sha256(baseline_bytes).hexdigest()
        == "0bda83e9e21f6db29cd827fe0ca1841dc5b12055265d3c25a461f4c0cf26b19a"
    )
    baseline = Document(BytesIO(baseline_bytes))

    def body_text(document):
        from docx.text.paragraph import Paragraph
        from docx.table import Table

        sequence = []
        for element in document.element.body:
            if element.tag == qn("w:p"):
                sequence.append(("p", Paragraph(element, document).text))
            elif element.tag == qn("w:tbl"):
                sequence.append(
                    (
                        "table",
                        [
                            [[p.text for p in cell.paragraphs] for cell in row.cells]
                            for row in Table(element, document).rows
                        ],
                    )
                )
        return sequence

    assert body_text(baseline) == body_text(doc), "First-delivery Word text differs"
    for table in doc.tables:
        assert table.rows[0]._tr.trPr.find(qn("w:tblHeader")) is not None
        for row in table.rows:
            assert row._tr.trPr.find(qn("w:cantSplit")) is not None
        assert all(
            p.paragraph_format.keep_with_next
            for c in table.rows[0].cells
            for p in c.paragraphs
        )
    assert not list(doc.styles["Title"].element.iter(qn("w:pBdr")))
    assert not list(doc.paragraphs[0]._p.iter(qn("w:pBdr")))
    for p in doc.paragraphs:
        for r in p.runs:
            assert str(r.font.color.rgb) == "000000"
    result = {
        "method": "Parse MD headings/paragraphs/table rows; reopen DOCX body XML in document order and parse HTML independently. Exact ordered text and cell comparison; no digest shortcut.",
        "semantic_blocks": len(expected),
        "tables": sum(k == "table" for k, _ in expected),
        "table_cells": sum(
            sum(len(row) for row in b) for k, b in expected if k == "table"
        ),
        "chinese_characters": len(re.findall(r"[\u4e00-\u9fff]", text)),
        "md_docx_html_sequence_equal": True,
        "mirror_equal": True,
        "all_sections_landscape": True,
        "first_delivery_md_bytes_equal": True,
        "first_delivery_word_ordered_text_equal": True,
        "repeat_headers_unsplit_rows_header_keep_next": True,
        "title_style_blue_border_removed": True,
        "baseline_word_ordered_text": body_text(baseline),
        "revised_word_ordered_text": body_text(doc),
        "all_paragraph_runs_black": True,
    }
    (ROOT / "qa").mkdir(exist_ok=True)
    (ROOT / "qa/revised").mkdir(exist_ok=True)
    (ROOT / "qa/revised/body-consistency.json").write_text(
        json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(
        json.dumps(
            {k: v for k, v in result.items() if not k.endswith("ordered_text")},
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    import sys

    if "--verify-only" not in sys.argv:
        generate()
    verify()
